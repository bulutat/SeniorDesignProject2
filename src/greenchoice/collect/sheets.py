"""Excel workbook for hand-collected data: template generator and CSV exporter.

People fill `data/raw/veri_toplama.xlsx`; `export_csvs` turns its two sheets into the
files the loaders expect (collect.pilot, collect.eprel_manual). Values are not validated
here, the loaders do that with row numbers.
"""

from __future__ import annotations

import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from greenchoice.collect.sheets_layout import (
    CATEGORY_LABELS,
    EPREL_COLUMNS,
    INSTRUCTIONS,
    PILOT_COLUMNS,
    Column,
)
from greenchoice.contracts import PRODUCT_COLUMNS

SHEET_HELP, SHEET_PRODUCTS, SHEET_EPREL = "Nasıl doldurulur", "Ürünler", "EPREL"
FIRST_DATA_ROW, LAST_DATA_ROW = 3, 300
REQUIRED_FILL = PatternFill("solid", fgColor="F4B183")
OPTIONAL_FILL = PatternFill("solid", fgColor="D9D9D9")


def format_sheet(ws, columns: list[Column]) -> None:
    """Widths and number formats only (no headers). Also repairs an existing, filled sheet:
    a date that carries a time part is too wide for its column and Excel shows ####."""
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = col.width
        if not col.number_format:
            continue
        for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
            cell = ws.cell(row=row, column=i)
            if isinstance(cell.value, datetime.datetime):
                cell.value = cell.value.date()
            cell.number_format = col.number_format


def _style_sheet(ws, columns: list[Column]) -> None:
    format_sheet(ws, columns)
    for i, col in enumerate(columns, start=1):
        label = ws.cell(row=1, column=i, value=col.label)
        label.font = Font(bold=True)
        label.fill = REQUIRED_FILL if col.required else OPTIONAL_FILL
        label.alignment = Alignment(wrap_text=True, vertical="center")
        hint = ws.cell(row=2, column=i, value=col.hint)
        hint.font = Font(italic=True, color="7F7F7F")
        if col.choices:
            dv = DataValidation(type="list", formula1='"' + ",".join(col.choices) + '"',
                                allow_blank=True)
            ws.add_data_validation(dv)
            letter = get_column_letter(i)
            dv.add(f"{letter}{FIRST_DATA_ROW}:{letter}{LAST_DATA_ROW}")
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = f"A{FIRST_DATA_ROW}"


def build_template(path: Path) -> None:
    wb = Workbook()
    help_ws = wb.active
    help_ws.title = SHEET_HELP
    help_ws.column_dimensions["A"].width = 110
    for row, line in enumerate(INSTRUCTIONS, start=1):
        help_ws.cell(row=row, column=1, value=line)
    _style_sheet(wb.create_sheet(SHEET_PRODUCTS), PILOT_COLUMNS)
    _style_sheet(wb.create_sheet(SHEET_EPREL), EPREL_COLUMNS)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        value = value.date()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _check_hint_row(ws, columns: list[Column]) -> None:
    """Row 2 holds hints. Typing a product there used to be skipped without any message."""
    hints = next(ws.iter_rows(min_row=2, max_row=2, max_col=len(columns), values_only=True), ())
    for col, cell in zip(columns, hints, strict=False):
        text = _cell_text(cell)
        if text and text != col.hint and text not in col.legacy_hints:
            raise ValueError(f"{ws.title} sheet, row 2 is the hint row but column {col.label!r} "
                             f"contains {text!r}. Data starts in row 3: move or clear this row.")


def _read_rows(ws, columns: list[Column]) -> list[dict]:
    _check_hint_row(ws, columns)
    rows = []
    for values in ws.iter_rows(min_row=FIRST_DATA_ROW, max_col=len(columns), values_only=True):
        row = {col.name: _cell_text(v) for col, v in zip(columns, values, strict=False)}
        if any(row.values()):
            rows.append(row)
    return rows


def _normalize_pilot(row: dict) -> dict:
    row["category"] = CATEGORY_LABELS.get(row["category"].lower(), row["category"])
    row["source"] = row["source"].lower()
    return row


def export_csvs(xlsx: Path, out_dir: Path) -> dict[str, Path]:
    wb = load_workbook(xlsx, data_only=True)
    pilot = [_normalize_pilot(r) for r in _read_rows(wb[SHEET_PRODUCTS], PILOT_COLUMNS)]
    eprel = _read_rows(wb[SHEET_EPREL], EPREL_COLUMNS)
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = {"pilot": out_dir / "pilot_products.csv", "eprel": out_dir / "eprel_manual.csv"}
    _write_csv(pilot, PRODUCT_COLUMNS, paths["pilot"])
    _write_csv(eprel, [c.name for c in EPREL_COLUMNS], paths["eprel"])
    return paths


def _write_csv(rows: list[dict], columns: list[str], path: Path) -> None:
    table = [{c: row.get(c, "") for c in columns} for row in rows]
    pd.DataFrame(table, columns=columns, dtype=str).to_csv(path, index=False, encoding="utf-8")
