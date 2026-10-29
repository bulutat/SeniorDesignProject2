import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from greenchoice.collect.eprel_manual import load_manual_eprel
from greenchoice.collect.pilot import load_pilot_list
from greenchoice.collect.sheets import build_template, export_csvs
from greenchoice.collect.sheets_layout import EPREL_COLUMNS, PILOT_COLUMNS
from greenchoice.contracts import PRODUCT_COLUMNS
from greenchoice.pipeline import main

URL = "https://eprel.ec.europa.eu/screen/product/refrigeratingappliances2019/1"


def fill(path, sheet, row, **values):
    columns = PILOT_COLUMNS if sheet == "Ürünler" else EPREL_COLUMNS
    wb = load_workbook(path)
    ws = wb[sheet]
    for i, col in enumerate(columns, start=1):
        if col.name in values:
            ws.cell(row=row, column=i, value=values[col.name])
    wb.save(path)


def product(**overrides):
    base = {"product_id": "ty-1", "source": "Trendyol", "url": "https://example.com/u/1",
            "category": "Buzdolabı", "brand": "Kuzey", "model_code": "KZ-1", "name": "Test",
            "price_tl": 12499.9, "energy_class": "C", "rating_count": 812.0}
    return base | overrides


@pytest.fixture
def template(tmp_path):
    path = tmp_path / "veri_toplama.xlsx"
    build_template(path)
    return path


def test_template_layout_is_readable(template):
    wb = load_workbook(template)
    assert wb.sheetnames == ["Nasıl doldurulur", "Ürünler", "EPREL"]
    ws = wb["Ürünler"]
    labels = [c.value for c in ws[1]][:len(PILOT_COLUMNS)]
    assert labels[0].endswith("*") and "Ürün" in labels[0]
    assert ws.freeze_panes == "A3"
    category_validations = [v for v in ws.data_validations.dataValidation
                            if "buzdolabı" in (v.formula1 or "")]
    assert category_validations


def test_export_converts_values_and_skips_blank_rows(template, tmp_path):
    fill(template, "Ürünler", 3, **product())
    fill(template, "Ürünler", 5, **product(product_id="ty-2", url="https://example.com/u/2",
                                           category="çamaşır makinesi"))
    out = export_csvs(template, tmp_path / "out")
    pilot = pd.read_csv(out["pilot"], dtype=str, keep_default_na=False)
    assert list(pilot.columns) == PRODUCT_COLUMNS
    assert pilot["product_id"].tolist() == ["ty-1", "ty-2"]
    assert pilot["category"].tolist() == ["refrigerator", "washing_machine"]
    assert pilot["source"].tolist() == ["trendyol", "trendyol"]
    assert pilot.loc[0, "price_tl"] == "12499.9"
    assert pilot.loc[0, "rating_count"] == "812"
    assert pilot.loc[0, "noise_db"] == ""


def test_exported_files_pass_the_real_loaders(template, tmp_path):
    fill(template, "Ürünler", 3, **product())
    fill(template, "EPREL", 3, model_code="KZ-1", energy_class="c", energy_kwh_per_year=182.5,
         eprel_url=URL, collected_on=datetime.date(2026, 10, 30))
    out = export_csvs(template, tmp_path / "out")
    assert load_pilot_list(out["pilot"]).loc[0, "price_tl"] == pytest.approx(12499.9)
    eprel = load_manual_eprel(out["eprel"])
    assert eprel.loc[0, "energy_class"] == "C"
    assert eprel.loc[0, "energy_kwh_per_year"] == pytest.approx(182.5)
    assert pd.read_csv(out["eprel"], dtype=str).loc[0, "collected_on"] == "2026-10-30"


@pytest.mark.parametrize("sheet", ["Ürünler", "EPREL"])
def test_data_typed_into_the_hint_row_is_an_error_not_silently_skipped(template, tmp_path, sheet):
    fill(template, sheet, 2, model_code="KZ-1", product_id="ty-1")
    with pytest.raises(ValueError, match=f"{sheet}.*row 2"):
        export_csvs(template, tmp_path / "out")


def test_cli_reports_hint_row_problem(template, tmp_path):
    fill(template, "Ürünler", 2, **product())
    with pytest.raises(SystemExit, match="row 2"):
        main(["import-sheets", "--xlsx", str(template), "--out-dir", str(tmp_path / "csv")])


def test_unknown_category_is_left_for_the_loader_to_reject(template, tmp_path):
    fill(template, "Ürünler", 3, **product(category="tost makinesi"))
    out = export_csvs(template, tmp_path / "out")
    with pytest.raises(ValueError, match="category"):
        load_pilot_list(out["pilot"])


def test_cli_make_and_import_sheets(tmp_path, capsys):
    xlsx = tmp_path / "sheet.xlsx"
    main(["make-sheets", "--out", str(xlsx)])
    assert xlsx.exists()
    fill(xlsx, "Ürünler", 3, **product())
    main(["import-sheets", "--xlsx", str(xlsx), "--out-dir", str(tmp_path / "csv")])
    assert (tmp_path / "csv" / "pilot_products.csv").exists()
    assert (tmp_path / "csv" / "eprel_manual.csv").exists()
    assert "pilot_products.csv" in capsys.readouterr().out


def test_cli_import_reports_row_problems_in_excel_terms(tmp_path):
    xlsx = tmp_path / "sheet.xlsx"
    build_template(xlsx)
    fill(xlsx, "Ürünler", 3, **product(model_code=None))
    with pytest.raises(SystemExit, match="Excel"):
        main(["import-sheets", "--xlsx", str(xlsx), "--out-dir", str(tmp_path / "csv")])
